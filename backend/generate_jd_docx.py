import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

ROOT_DIR = Path(r"e:\Resume screnner")

JD_LIST = [
    {
        "filename": "JD1_Senior_AI_LLM_Engineer.docx",
        "title": "Senior AI & LLM Systems Engineer",
        "department": "Artificial Intelligence & Core Systems",
        "employment_type": "Full-time",
        "experience": "5+ years of professional engineering experience",
        "location": "Remote / Hybrid",
        "about": (
            "We are seeking an experienced Senior AI & LLM Systems Engineer to lead the architecture, development, "
            "and deployment of enterprise Generative AI pipelines and intelligent multi-agent systems. In this role, "
            "you will bridge state-of-the-art Large Language Models (LLMs) with high-performance backend infrastructure, "
            "focusing on Retrieval-Augmented Generation (RAG), vector embeddings, and real-time inference services."
        ),
        "responsibilities": [
            "Design, build, and optimize scalable RAG pipelines using modern vector search engines and embedding models.",
            "Fine-tune and evaluate open-source and proprietary LLM models for high domain accuracy and minimal latency.",
            "Build resilient, low-latency API microservices with Python and FastAPI containerized with Docker on Kubernetes.",
            "Implement robust evaluation frameworks, guardrails, and automated verification tests to mitigate hallucinations.",
            "Collaborate with product and software engineering teams to integrate AI capabilities into production applications."
        ],
        "must_have": [
            "At least 5 years of professional software engineering and machine learning experience.",
            "Strong proficiency in Python, PyTorch, and modern machine learning frameworks.",
            "Deep hands-on experience building Retrieval-Augmented Generation (RAG) systems using LangChain or LlamaIndex.",
            "Experience with vector databases such as ChromaDB, Pinecone, or FAISS.",
            "Proven experience building and deploying production REST APIs using FastAPI or Flask.",
            "Working knowledge of containerization and deployment with Docker, Kubernetes, and AWS cloud services."
        ],
        "nice_to_have": [
            "Experience with Hugging Face Transformers and fine-tuning open-source LLMs (Llama, Mistral, Qwen).",
            "Familiarity with MLOps tracking tools (MLflow, Weights & Biases) and CI/CD pipelines.",
            "Solid understanding of SQL, PostgreSQL, and vector similarity search indexing algorithms (HNSW, IVFFlat).",
            "Background in prompt engineering, agentic workflows, and semantic re-ranking."
        ]
    },
    {
        "filename": "JD2_Senior_Full_Stack_Cloud_Engineer.docx",
        "title": "Senior Full Stack & Cloud Systems Engineer",
        "department": "Product Engineering",
        "employment_type": "Full-time",
        "experience": "5+ years of professional software development experience",
        "location": "Remote / Hybrid",
        "about": (
            "We are looking for a versatile Senior Full Stack & Cloud Systems Engineer to architect and build resilient "
            "web applications and microservices. You will work across modern responsive frontend interfaces, "
            "high-throughput backend APIs, and automated cloud deployments on AWS and Kubernetes."
        ),
        "responsibilities": [
            "Architect and develop modern responsive web applications using React, TypeScript, and modern CSS/HTML.",
            "Design and scale distributed backend microservices and RESTful APIs using Python (FastAPI/Django) or Node.js.",
            "Model, optimize, and maintain relational databases (PostgreSQL, MySQL) and caching systems (Redis).",
            "Build and maintain continuous integration and continuous delivery (CI/CD) pipelines with Docker, Kubernetes, and AWS.",
            "Collaborate with cross-functional teams in an Agile/Scrum environment to deliver high-quality, tested software."
        ],
        "must_have": [
            "5+ years of professional full stack or backend software engineering experience.",
            "Strong proficiency in modern JavaScript/TypeScript and frontend libraries such as React.",
            "Solid backend development experience using Python (FastAPI, Django) or Node.js.",
            "Strong database knowledge with PostgreSQL, MySQL, and relational schema optimization.",
            "Hands-on experience with containerization using Docker and orchestrating services with Kubernetes or AWS.",
            "Firm commitment to clean code, Git version control, and automated testing (CI/CD)."
        ],
        "nice_to_have": [
            "Experience with Go (Golang) and high-throughput event streaming (Apache Kafka).",
            "Familiarity with Infrastructure as Code (IaC) using Terraform or Ansible.",
            "Experience with GraphQL APIs, Redis caching, or microservices migration.",
            "Experience building responsive design systems and accessible web interfaces."
        ]
    }
]


def build_docx(jd_data: dict, out_path: Path):
    doc = docx.Document()

    # Set normal margins (1 inch)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # 1. Main Title
    title_p = doc.add_paragraph()
    title_run = title_p.add_run(jd_data["title"])
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)  # Navy Blue
    title_p.paragraph_format.space_after = Pt(4)

    # 2. Metadata subtitle
    meta_p = doc.add_paragraph()
    meta_run = meta_p.add_run(
        f"Department: {jd_data['department']}  |  Employment Type: {jd_data['employment_type']}  |  Location: {jd_data['location']}\n"
        f"Experience Required: {jd_data['experience']}"
    )
    meta_run.font.name = "Calibri"
    meta_run.font.size = Pt(10)
    meta_run.font.italic = True
    meta_run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)  # Slate Gray
    meta_p.paragraph_format.space_after = Pt(14)

    # Helper for adding section headings
    def add_heading(text: str):
        h = doc.add_paragraph()
        run = h.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)  # Dark Navy
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)

    # Helper for body paragraphs
    def add_body(text: str):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(6)

    # Helper for bullet points
    def add_bullets(items: list[str]):
        for item in items:
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(item)
            run.font.name = "Calibri"
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(3)

    # 3. About the Role
    add_heading("About the Role")
    add_body(jd_data["about"])

    # 4. Key Responsibilities
    add_heading("Key Responsibilities")
    add_bullets(jd_data["responsibilities"])

    # 5. Must-Have Requirements
    add_heading("Must-Have Requirements")
    add_bullets(jd_data["must_have"])

    # 6. Nice-to-Have / Preferred Qualifications
    add_heading("Nice-to-Have / Preferred Qualifications")
    add_bullets(jd_data["nice_to_have"])

    doc.save(str(out_path))
    print(f"Generated DOCX: {out_path.name}")


def main():
    for item in JD_LIST:
        out_file = ROOT_DIR / item["filename"]
        build_docx(item, out_file)


if __name__ == "__main__":
    main()
